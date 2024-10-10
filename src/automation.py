import pandas as pd
import pywhatkit as pwk
import time
from docx import Document
from datetime import datetime
import os
import openpyxl

# Load Excel contact list and message template
def load_files(contacts_path, message_template_path):
    df = pd.read_excel(contacts_path)
    doc = Document(message_template_path)
    message = '\n'.join([paragraph.text for paragraph in doc.paragraphs])
    return df, message

# Update Excel file with message status
def update_status(status_file, name, email, phone_number, status):
    timestamp = datetime.now().strftime('%d-%m-%Y %H:%M:%S')
    new_status = pd.DataFrame([[name, email, phone_number, status, timestamp]],
                              columns=['Name', 'Email', 'PhoneNumber', 'Status', 'Timestamp'])

    # Check if the status file exists and create or append data accordingly
    if os.path.exists(status_file):
        with pd.ExcelWriter(status_file, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
            startrow = writer.sheets['MessageStatus'].max_row
            new_status.to_excel(writer, sheet_name='MessageStatus', index=False, startrow=startrow, header=False)
    else:
        with pd.ExcelWriter(status_file, engine='openpyxl', mode='w') as writer:
            new_status.to_excel(writer, sheet_name='MessageStatus', index=False)

# Send WhatsApp messages
def send_messages(df, message, status_file, stop_flag):
    success_count = 0
    failure_count = 0

    for index, row in df.iterrows():
        if stop_flag:
            break

        name = row['Name']
        email = row['Email']
        phone_number = str(row['PhoneNumber']).strip()

        if not phone_number.startswith('+91'):
            phone_number = '+91' + phone_number

        try:
            pwk.sendwhatmsg_instantly(phone_number, message, wait_time=10, tab_close=True, close_time=3)
            update_status(status_file, name, email, phone_number, "Success")
            success_count += 1
            time.sleep(5)

        except Exception as e:
            error_message = str(e)
            update_status(status_file, name, email, phone_number, f"Failed: {error_message}")
            failure_count += 1

    return success_count, failure_count
